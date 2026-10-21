from io import BytesIO
from zipfile import ZipFile

from docx import Document

from routes.docx_report import build_docx_report


def _docx(report, **kwargs):
    payload = build_docx_report(report=report, title=kwargs.pop("title", "Test Report"), **kwargs)
    assert payload[:2] == b"PK"
    assert ZipFile(BytesIO(payload)).testzip() is None
    return Document(BytesIO(payload)), payload


def _all_text(document):
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def test_docx_export_has_editable_text_tables_chart_data_and_no_control_markers():
    report = """## Renewable capacity
Supported capacity reached 12.5 GW [S1, S2].

| Sector | Capacity |
|---|---:|
| Solar | 10.0 GW |
| Wind | 2.5 GW |

[RISK_MATRIX]

[CHART_1]
"""
    charts = [{
        "type": "bar",
        "title": "Renewable capacity by sector",
        "unit": "GW",
        "series": [{"name": "Installed capacity", "data": [
            {"label": "Solar", "value": 10.0},
            {"label": "Wind", "value": 2.5},
        ]}],
    }]
    sources = [
        {"id": "S1", "title": "Official source", "url": "https://example.com/official", "publisher": "Example", "page_fetch_status": "ok"},
        {"id": "S2", "title": "Supporting source", "url": "https://example.com/support"},
    ]
    document, payload = _docx(report, question="Energy outlook", key_stats=[{"label": "Capacity", "value": "12.5 GW"}], charts=charts, sources=sources)
    body = _all_text(document)
    assert "Renewable capacity" in body
    assert "Solar" in body and "10.0 GW" in body
    assert "Renewable capacity by sector" in body
    assert "12.5 GW" in body
    assert "[RISK_MATRIX]" not in body
    assert "[CHART_1]" not in body
    assert "Official source" in body and "Supporting source" in body
    assert "Risk matrix data was not supplied" in body
    assert any(rel.reltype.endswith("/hyperlink") and rel.is_external for rel in document.part.rels.values())


def test_docx_does_not_force_executive_summary_or_conclusion():
    document, _ = _docx("## Findings\nOnly evidence-supported prose.", question="Test")
    body = _all_text(document)
    assert "Findings" in body
    assert "Executive Summary" not in body
    assert "Conclusion" not in body


def test_docx_preserves_unplaced_charts_and_marks_uploaded_document_provenance():
    charts = [{"type": "line", "title": "Trend", "series": [{"name": "Value", "data": [
        {"label": "2024", "value": 4}, {"label": "2025", "value": 7},
    ]}]}]
    document, _ = _docx(
        "## Outlook\nThe future depends on execution.",
        charts=charts,
        source_documents=[{"name": "uploaded.xlsx", "file_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "text": "not copied into output"}],
    )
    body = _all_text(document)
    assert "Editable chart data" in body
    assert "Trend" in body and "2024" in body and "2025" in body
    assert "Uploaded / RAG Documents" in body
    assert "uploaded.xlsx" in body
    assert "not copied into output" not in body


def test_docx_can_build_table_from_chart_columns_and_rows_without_fabricating_values():
    charts = [{"type": "table", "title": "Source table", "columns": ["Year", "Value"], "rows": [["2024", "5"], ["2025", "7"]]}]
    document, _ = _docx("## Results", charts=charts)
    body = _all_text(document)
    assert "Source table" in body
    assert "2024" in body and "5" in body and "2025" in body and "7" in body
    assert "No structured observations were supplied" not in body


def test_supported_chart_is_a_native_word_chart_with_embedded_editable_workbook():
    from openpyxl import load_workbook

    chart = {"type": "bar", "title": "Capacity", "series": [{"name": "GW", "data": [
        {"label": "Solar", "value": 10}, {"label": "Wind", "value": 5},
    ]}]}
    _, payload = _docx("## Results\n[CHART_1]", charts=[chart])
    with ZipFile(BytesIO(payload)) as archive:
        names = archive.namelist()
        chart_parts = [name for name in names if name.startswith("word/charts/chart") and name.endswith(".xml")]
        workbook_parts = [name for name in names if name.startswith("word/embeddings/") and name.endswith(".xlsx")]
        assert chart_parts, "supported charts should be real Word chart objects, not only pictures"
        assert workbook_parts, "editable chart must include its worksheet data"
        chart_xml = archive.read(chart_parts[0]).decode("utf-8")
        assert "Capacity" in chart_xml and "Solar" in chart_xml and "Wind" in chart_xml
        workbook = load_workbook(BytesIO(archive.read(workbook_parts[0])), data_only=True)
        sheet = workbook["Data"]
        assert sheet["A2"].value == "Solar" and sheet["B2"].value == 10
        assert sheet["A3"].value == "Wind" and sheet["B3"].value == 5


def test_hydrated_images_are_embedded_and_image_markers_never_leak():
    import base64
    import io
    from PIL import Image

    image_buffer = io.BytesIO()
    Image.new("RGB", (3, 3), (20, 120, 80)).save(image_buffer, format="PNG")
    encoded = base64.b64encode(image_buffer.getvalue()).decode("ascii")
    document, payload = _docx(
        "## Visuals\n[IMAGE_1]\n[WEB_IMG_2]",
        images=[
            {"data": "data:image/png;base64," + encoded, "caption": "Evidence image"},
            {"data": "data:image/png;base64," + encoded, "caption": "Evidence image alias"},
        ],
    )
    assert len(document.inline_shapes) >= 1
    assert "Evidence image" in _all_text(document)
    assert "Evidence image alias" in _all_text(document)
    with ZipFile(BytesIO(payload)) as archive:
        assert any(name.startswith("word/media/image") for name in archive.namelist())
    assert "[IMAGE_1]" not in _all_text(document)
    assert "[WEB_IMG_2]" not in _all_text(document)


def test_long_research_question_is_clamped_to_core_properties_limit():
    question = "Research question " + ("with additional context " * 30)
    document, _ = _docx("## Findings\nEvidence-backed narrative [S1].", question=question)

    assert len(document.core_properties.subject) <= 255
    assert question[:500] in _all_text(document)
